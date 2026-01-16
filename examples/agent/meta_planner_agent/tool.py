# -*- coding: utf-8 -*-
"""The tool functions used in the planner example."""
import asyncio
import json
import os
from collections import OrderedDict
from typing import AsyncGenerator

from pydantic import BaseModel, Field

from agentscope.agent import ReActAgent
from agentscope.formatter import DashScopeChatFormatter
from agentscope.mcp import HttpStatelessClient, StdIOStatefulClient
from agentscope.message import Msg, TextBlock
from agentscope.model import DashScopeChatModel
from agentscope.pipeline import stream_printing_messages
from agentscope.tool import (
    ToolResponse,
    Toolkit,
    write_text_file,
    insert_text_file,
    view_text_file,
)

try:
    import pandas as pd  # noqa: F401

    PANDAS_AVAILABLE = True
except ImportError:
    PANDAS_AVAILABLE = False

try:
    import openpyxl  # noqa: F401

    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


class ResultModel(BaseModel):
    """
    The result model used for the sub worker to summarize the task result.
    """

    success: bool = Field(
        description="Whether the task was successful or not.",
    )
    message: str = Field(
        description=(
            "The specific task result, should include necessary details, "
            "e.g. the file path if any file is generated, the deviation, "
            "and the error message if any."
        ),
    )


async def read_excel_file(
    file_path: str,
    sheet_name: str | None = None,
    max_rows: int = 1000,
    include_index: bool = False,
) -> ToolResponse:
    """Read and analyze an Excel (.xlsx, .xls) file.

    This function can read Excel files and return their contents in a readable
    text format. It supports multiple sheets and can limit the number of rows
    returned to avoid overwhelming the context.

    Args:
        file_path (`str`):
            The path to the Excel file to read. Can be an absolute or relative
            path.
        sheet_name (`str | None`, optional):
            The name of the sheet to read. If None, reads the first sheet.
            If "all", reads all sheets.
        max_rows (`int`, defaults to 1000):
            Maximum number of rows to return per sheet. This helps control
            memory usage for large files.
        include_index (`bool`, defaults to False):
            Whether to include row indices in the output.

    Returns:
        `ToolResponse`:
            A ToolResponse containing the Excel file contents as text blocks.
    """
    import os

    if not os.path.exists(file_path):
        return ToolResponse(
            content=[
                TextBlock(
                    type="text",
                    text=f"Error: File '{file_path}' does not exist.",
                ),
            ],
        )

    try:
        if PANDAS_AVAILABLE:
            # Use pandas for reading Excel files
            if sheet_name == "all":
                # Read all sheets
                excel_file = pd.ExcelFile(file_path)
                sheets_data = {}
                for sheet in excel_file.sheet_names:
                    df = pd.read_excel(excel_file, sheet_name=sheet, nrows=max_rows)
                    sheets_data[sheet] = df

                content_blocks = [
                    TextBlock(
                        type="text",
                        text=f"Excel file '{file_path}' contains {len(sheets_data)} sheet(s):\n",
                    ),
                ]

                for sheet_name_key, df in sheets_data.items():
                    content_blocks.append(
                        TextBlock(
                            type="text",
                            text=f"\n--- Sheet: {sheet_name_key} ({len(df)} rows) ---\n",
                        ),
                    )
                    # Convert DataFrame to string representation
                    df_str = df.to_string(index=include_index, max_rows=max_rows)
                    content_blocks.append(
                        TextBlock(
                            type="text",
                            text=df_str,
                        ),
                    )
            else:
                # Read specific sheet or first sheet
                result = pd.read_excel(
                    file_path,
                    sheet_name=sheet_name,
                    nrows=max_rows,
                )

                # Handle case where read_excel returns a dict (multiple sheets)
                if isinstance(result, dict):
                    # If dict is returned, use the first sheet or specified sheet
                    if sheet_name and sheet_name in result:
                        df = result[sheet_name]
                    else:
                        # Get first sheet from dict
                        first_sheet = list(result.keys())[0]
                        df = result[first_sheet]
                        sheet_name = first_sheet
                else:
                    # Single DataFrame returned
                    df = result

                content_blocks = [
                    TextBlock(
                        type="text",
                        text=f"Excel file '{file_path}' (Sheet: {sheet_name or 'first'}, {len(df)} rows):\n",
                    ),
                ]

                # Convert DataFrame to string representation
                df_str = df.to_string(index=include_index, max_rows=max_rows)
                content_blocks.append(
                    TextBlock(
                        type="text",
                        text=df_str,
                    ),
                )

            return ToolResponse(content=content_blocks)

        elif OPENPYXL_AVAILABLE:
            # Use openpyxl as fallback
            from openpyxl import load_workbook

            wb = load_workbook(file_path, data_only=True)

            if sheet_name == "all":
                sheets_data = {}
                for sheet in wb.sheetnames:
                    ws = wb[sheet]
                    rows = []
                    for row_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
                        if row_idx > max_rows:
                            break
                        rows.append(row)
                    sheets_data[sheet] = rows

                content_blocks = [
                    TextBlock(
                        type="text",
                        text=f"Excel file '{file_path}' contains {len(sheets_data)} sheet(s):\n",
                    ),
                ]

                for sheet_name_key, rows in sheets_data.items():
                    content_blocks.append(
                        TextBlock(
                            type="text",
                            text=f"\n--- Sheet: {sheet_name_key} ({len(rows)} rows) ---\n",
                        ),
                    )
                    # Convert rows to text
                    rows_text = "\n".join(
                        ["\t".join([str(cell) if cell is not None else "" for cell in row]) for row in rows]
                    )
                    content_blocks.append(
                        TextBlock(
                            type="text",
                            text=rows_text,
                        ),
                    )
            else:
                # Read specific sheet or first sheet
                ws = wb[sheet_name] if sheet_name else wb.active
                rows = []
                for row_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
                    if row_idx > max_rows:
                        break
                    rows.append(row)

                content_blocks = [
                    TextBlock(
                        type="text",
                        text=f"Excel file '{file_path}' (Sheet: {ws.title}, {len(rows)} rows):\n",
                    ),
                ]

                # Convert rows to text
                rows_text = "\n".join(
                    ["\t".join([str(cell) if cell is not None else "" for cell in row]) for row in rows]
                )
                content_blocks.append(
                    TextBlock(
                        type="text",
                        text=rows_text,
                    ),
                )

            return ToolResponse(content=content_blocks)

        else:
            return ToolResponse(
                content=[
                    TextBlock(
                        type="text",
                        text=(
                            "Error: Neither pandas nor openpyxl is available. "
                            "Please install one of them to read Excel files:\n"
                            "  pip install pandas\n"
                            "  or\n"
                            "  pip install openpyxl"
                        ),
                    ),
                ],
            )

    except Exception as e:
        return ToolResponse(
            content=[
                TextBlock(
                    type="text",
                    text=f"Error reading Excel file '{file_path}': {str(e)}",
                ),
            ],
        )


def _convert_to_text_block(msgs: list[Msg]) -> list[TextBlock]:
    # Collect all the content blocks
    blocks: list = []
    # Convert tool_use block into text block for streaming tool response
    for _ in msgs:
        for block in _.get_content_blocks():
            if block["type"] == "text":
                blocks.append(block)

            elif block["type"] == "tool_use":
                blocks.append(
                    TextBlock(
                        type="text",
                        text=f"Calling tool {block['name']} ...",
                    ),
                )

    return blocks


async def create_worker(
    task_description: str,
) -> AsyncGenerator[ToolResponse, None]:
    """Create a sub-worker agent to finish the given task.

    The sub-worker agent is equipped with various tools to handle different types of tasks:
    - **Browser tools**: Web browsing, page navigation, content extraction from websites
      (e.g., accessing arXiv.org, reading web pages, extracting information)
    - **GitHub tools**: Repository search and code file retrieval
    - **Map tools**: Geocoding, routing, and place search (if GAODE_API_KEY is set)
    - **File tools**: Read, write, and view text files
    - **Excel tools**: Read and analyze Excel files

    Use this tool when you need to:
    - Access websites or web pages (e.g., arXiv.org papers, online articles)
    - Search GitHub repositories or retrieve code files
    - Read or analyze files
    - Perform web-based research or data extraction

    Args:
        task_description (`str`):
            The description of the task to be done by the sub-worker, should
            contain all the necessary information. Be specific about:
            - URLs to visit (if web access is needed)
            - Files to read or analyze
            - Specific information to extract or find
            - Expected output format

    Returns:
        `AsyncGenerator[ToolResponse, None]`:
            An async generator yielding ToolResponse objects with the execution
            process and final results.
    """
    toolkit = Toolkit()

    # Gaode MCP client
    if os.getenv("GAODE_API_KEY"):
        toolkit.create_tool_group(
            group_name="amap_tools",
            description="Map-related tools, including geocoding, routing, and " "place search.",
        )
        client = HttpStatelessClient(
            name="amap_mcp",
            transport="streamable_http",
            url=f"https://mcp.amap.com/mcp?key={os.environ['GAODE_API_KEY']}",
        )
        await toolkit.register_mcp_client(client, group_name="amap_tools")
    else:
        print(
            "Warning: GAODE_API_KEY not set in environment, skipping Gaode " "MCP client registration.",
        )

    # Browser MCP client
    toolkit.create_tool_group(
        group_name="browser_tools",
        description="Web browsing related tools.",
    )
    browser_client = StdIOStatefulClient(
        name="playwright-mcp",
        command="npx",
        args=["@playwright/mcp@latest"],
    )
    await browser_client.connect()
    await toolkit.register_mcp_client(
        browser_client,
        group_name="browser_tools",
    )

    # GitHub MCP client
    if os.getenv("GITHUB_TOKEN"):
        toolkit.create_tool_group(
            group_name="github_tools",
            description="GitHub related tools, including repository " "search and code file retrieval.",
        )
        github_client = HttpStatelessClient(
            name="github",
            transport="streamable_http",
            url="https://api.githubcopilot.com/mcp/",
            headers={"Authorization": f"Bearer {os.getenv('GITHUB_TOKEN')}"},
        )
        await toolkit.register_mcp_client(
            github_client,
            group_name="github_tools",
        )

    else:
        print(
            "Warning: GITHUB_TOKEN not set in environment, skipping GitHub " "MCP client registration.",
        )

    # Basic read/write tools
    toolkit.register_tool_function(write_text_file)
    toolkit.register_tool_function(insert_text_file)
    toolkit.register_tool_function(view_text_file)

    # Excel file reading tool
    toolkit.register_tool_function(read_excel_file)

    # Create a new sub-agent to finish the given task
    sub_agent = ReActAgent(
        name="Worker",
        sys_prompt=f"""You're an agent named Worker.

## Your Target
Your target is to finish the given task with your tools.

## IMPORTANT
You MUST use the `{ReActAgent.finish_function_name}` to generate the final answer after finishing the task.
""",  # noqa: E501  # pylint: disable=C0301
        model=DashScopeChatModel(
            model_name="qwen3-max",
            api_key=os.environ["DASHSCOPE_API_KEY"],
        ),
        enable_meta_tool=True,
        formatter=DashScopeChatFormatter(),
        toolkit=toolkit,
        max_iters=20,
    )

    # disable the console output of the sub-agent
    sub_agent.set_console_output_enabled(False)

    # Collect the execution process content
    msgs = OrderedDict()

    # Wrap the sub-agent in a coroutine task to obtain the final
    # structured output
    result = []

    async def call_sub_agent() -> None:
        msg_res = await sub_agent(
            Msg(
                "user",
                content=task_description,
                role="user",
            ),
            structured_model=ResultModel,
        )
        result.append(msg_res)

    # Use stream_printing_message to get the streaming response as the
    # sub-agent works
    async for msg, _ in stream_printing_messages(
        agents=[sub_agent],
        coroutine_task=call_sub_agent(),
    ):
        msgs[msg.id] = msg

        # Collect all the content blocks
        yield ToolResponse(
            content=_convert_to_text_block(
                list(msgs.values()),
            ),
            stream=True,
            is_last=False,
        )

        # Expose the interruption signal to the caller
        if msg.metadata and msg.metadata.get("_is_interrupted", False):
            raise asyncio.CancelledError()

    # Obtain the last message from the coroutine task
    if result:
        yield ToolResponse(
            content=[
                *_convert_to_text_block(
                    list(msgs.values()),
                ),
                TextBlock(
                    type="text",
                    text=json.dumps(
                        result[0].metadata,
                        indent=2,
                        ensure_ascii=False,
                    ),
                ),
            ],
            stream=True,
            is_last=True,
        )

    await browser_client.close()
